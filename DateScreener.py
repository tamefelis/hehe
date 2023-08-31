import streamlit as st
import pandas as pd
from docx import Document
import shutil
from dateutil import parser
import os 

class DateScreener:
    def __init__(self, uploaded_file):
        xls = pd.ExcelFile(uploaded_file)
        self.sheet_name = st.selectbox('Choose a worksheet:', xls.sheet_names)
        self.df = xls.parse(self.sheet_name)

    @staticmethod
    def to_datetime(df, columns):
        for col in columns:
            df[col] = pd.to_datetime(df[col].astype(str).str.strip(), format='%d/%m/%Y', errors='coerce')
        return df

    @staticmethod
    def filter_dates(df, column_name, start_date, end_date):
        start_datetime = pd.Timestamp(start_date)
        end_datetime = pd.Timestamp(end_date)
        return df[(df[column_name] >= start_datetime) & (df[column_name] <= end_datetime)]

    def preprocess_screening_visit(self, start_date, end_date):
        df = self.to_datetime(self.df, ['Date for Screening'])
        df['Contact Number'] = df['Contact Number'].fillna(0).astype(int).astype(str)
        return self.filter_dates(df, 'Date for Screening', start_date, end_date)[['Screening ID', 'Date for Screening', 'Start time', 'Contact Number']]


    def actual_visit(self, start_date, end_date):
        visit_cols = ['Visit 1', 'Visit 2', 'Visit 3', 'Visit 4']
        time_cols = ['V1_time', 'V2_time', 'V3_time', 'V4_time']
        df = self.to_datetime(self.df, visit_cols)
        df['V_time'] = df[time_cols].apply(lambda x: ', '.join(x.dropna().astype(str)), axis=1)
        
        # Initialize an empty DataFrame to store the result
        df_continuous = pd.DataFrame()
        
        # Iterate through visit columns
        for col in visit_cols:
            df['Contact Number'] = df['Contact Number'].fillna(0).astype(int).astype(str)
            temp_df = df.dropna(subset=[col])[[col, 'Study ID', 'V_time','Contact Number']].rename(columns={col: 'Date'})
            temp_df['Visit_Type'] = col  # Add the visit type as a new column
            df_continuous = pd.concat([df_continuous, temp_df])
    
        df_continuous.sort_values('Date', inplace=True)
        df_continuous['Cumulative_Count'] = range(1, len(df_continuous) + 1)
        df_continuous = self.filter_dates(df_continuous, 'Date', start_date, end_date)

        return df_continuous[['Study ID', 'Date', 'V_time', 'Visit_Type', 'Contact Number']]
    # Include the 'Visit_Type' column in the return
    
    def display_and_return_df(self, df):
        st.write(df)
        st.write(f"Total count of rows: {len(df)}")
        return df

#    @staticmethod
#    def edit_word_document(template_path, study_id, visit_date, v_time, visit_type):
#    # Open the template
#        doc = Document(template_path)
#        # Determine the replacement for the placeholder
#        visit_number = visit_type.split(' ')[-1]
#        if visit_type == 'Visit 2':
#            return None  # Do not edit the document for Visit 2#

#        replacement_id = f'participant_ID_V{visit_number}'
#        replacement_date = visit_date.strftime('%d/%m/%Y')#

#        # Parse the time into a consistent format
#        parsed_time = parser.parse(v_time)
#        replacement_time = parsed_time.strftime('%H:%M')#

#        # Iterate through the paragraphs and replace the placeholders
#        for paragraph in doc.paragraphs:
#            for run in paragraph.runs:
#                run.text = run.text.replace('participant_ID_V%', replacement_id)
#                run.text = run.text.replace('DD/MM/YYYY', replacement_date)
#                run.text = run.text.replace('HH:MM', replacement_time)#

#        # Save as a new file
#        new_file_path = f"C:\\Users\\yuenchl\\Desktop\\AH DDI\\edited_template_{study_id}_{visit_number}.docx"
#        doc.save(new_file_path)
#        # Log the change
#        log_message = f"Edited document for Study ID: {study_id}, Visit Type: {visit_type}\n"
#        with open('log.txt', 'a') as log_file:
#            log_file.write(log_message)
#        
#        # Log the change on Streamlit
#        st.write(log_message)#

#        return new_file_path

def main():
    st.title('Date Screener')
    uploaded_file = st.file_uploader("Upload an Excel file", type=["xlsx"])

    if uploaded_file:
        screener = DateScreener(uploaded_file)
        start_date = st.date_input("Select a start date")
        end_date = st.date_input("Select an end date")
        if screener.sheet_name == 'Calendared Screening Visit':
            df = screener.preprocess_screening_visit(start_date, end_date)
            st.write(df)
            st.write(f"Total count of rows: {len(df)}")  # Displaying the count of rows
        elif screener.sheet_name == 'Calendared Study Visit':
            df = screener.actual_visit(start_date, end_date)
            st.write(df)
            st.write(f"Total count of rows: {len(df)}")  # Displaying the count of rows

if __name__ == "__main__":
    main()


    #template_file = st.file_uploader("Upload a Word template", type=["docx"])

## Button to activate Word editing
#    if st.button('Edit Word Document'):
#        if uploaded_file and template_file:
#            # Temporarily save the uploaded Word file
#            template_path = 'template.docx'
#            with open(template_path, 'wb') as f:
#                shutil.copyfileobj(template_file, f)
#
#            # Process the DataFrame and edit the Word document
#            if screener.sheet_name == 'Calendared Study Visit':
#                df_actual_visit = screener.actual_visit(start_date, end_date)
#                for index, row in df_actual_visit.iterrows():
#                    study_id = row['Study ID']
#                    visit_date = row['Date']
#                    v_time = row['V_time']
#                    visit_type = row['Visit_Type']
#                    new_file_path = screener.edit_word_document(template_path, study_id, visit_date, v_time, visit_type)
#                    if new_file_path:
#                        file_size = os.path.getsize(new_file_path)
#                        with open(new_file_path, 'rb') as f:
#                            st.download_button(
#                                f"Download the edited file for Study ID: {study_id}, Visit Type: {visit_type}",
#                                f.read(),
#                                file_size=file_size,
#                                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
#                                file_name=f"edited_template.docx"
